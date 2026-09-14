# -*- coding: utf-8 -*-
"""授课PPT生成器 —— 按《PPT模版.pptx》配色/风格/版式
配色: 深蓝#0545E8 藏青#010328 橙#F5670F 辅蓝#0475B6 白底
字体: 微软雅黑(中文)/Calibri(西文)  版式: 16:9
每个PPT = 一次课(2课时, 80分钟), 约10-12页
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

BLUE = RGBColor(0x05, 0x45, 0xE8)
NAVY = RGBColor(0x01, 0x03, 0x28)
ORANGE = RGBColor(0xF5, 0x67, 0x0F)
CBLUE = RGBColor(0x04, 0x75, 0xB6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x44, 0x54, 0x6A)
SW, SH = Emu(12192000), Emu(6858000)      # 16:9
FONT = "微软雅黑"


def _set(run, size, color=NAVY, bold=False, font=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font


def box(slide, x, y, w, h, fill=None, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    sh.shadow.inherit = False
    return sh


def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6, line_spacing=1.1):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for para in runs:                      # para: list[(txt,size,color,bold)]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for (t, sz, c, b) in para:
            r = p.add_run()
            r.text = t
            _set(r, sz, c, b)
    return tb


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SW, SH

    def add(self):
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])  # blank

    # ---------- 1 封面 ----------
    def cover(self, course, title, sub, meta, dept=None, illus=None):
        s = self.add()
        box(s, 0, 0, SW, SH * 26 // 100, NAVY)
        box(s, 0, SH * 26 // 100, SW, Emu(60000), ORANGE)
        rows = [[(course, 22, WHITE, False)]]
        if dept:
            rows.append([("◆ 重庆两江智造集团 · " + dept, 13,
                          RGBColor(0xBB, 0xC6, 0xE6), False)])
        text(s, Emu(700000), Emu(250000), SW - Emu(1400000), Emu(1100000),
             rows, PP_ALIGN.LEFT)
        tw = SW - Emu(6200000) if illus else SW - Emu(1400000)
        text(s, Emu(700000), SH * 33 // 100, tw, Emu(2200000),
             [[(title, 34, NAVY, True)]])
        text(s, Emu(700000), SH * 55 // 100, tw, Emu(800000),
             [[(sub, 18, BLUE, False)], [(meta, 14, GRAY, False)]])
        if illus:
            try:
                s.shapes.add_picture(illus, Emu(7500000), Emu(2300000),
                                     Emu(4100000), Emu(3075000))
            except Exception:
                pass
        box(s, Emu(700000), SH * 82 // 100, Emu(1500000), Emu(80000), ORANGE)
        text(s, Emu(700000), SH * 84 // 100, SW - Emu(1400000), Emu(500000),
             [[("重庆现代制造职业学院  智能控制学院 · 工业互联网应用专业", 12, GRAY, False)]])

    # ---------- 2 课堂流程(时间盒) ----------
    def flow(self, rows, note="每课时40分钟 · 本次课2课时共80分钟", note_emu=None):
        s = self.add()
        self._header(s, "本次课流程（80分钟时间盒）")
        y = Emu(1900000)
        for i, (seg, t, act) in enumerate(rows):
            box(s, Emu(800000), y, Emu(10000000), Emu(680000),
                fill=WHITE if i % 2 else RGBColor(0xF3, 0xF6, 0xFD))
            box(s, Emu(800000), y, Emu(1700000), Emu(680000),
                fill=ORANGE if seg in ("决策", "评价") else BLUE)
            text(s, Emu(830000), y + Emu(100000), Emu(1600000), Emu(480000),
                 [[(seg, 16, WHITE, True)]], PP_ALIGN.CENTER)
            text(s, Emu(2700000), y + Emu(100000), Emu(1500000), Emu(480000),
                 [[(t, 16, NAVY, True)]])
            text(s, Emu(4300000), y + Emu(100000), Emu(6400000), Emu(520000),
                 [[(act, 14, GRAY, False)]])
            y += Emu(790000)
        text(s, Emu(800000), Emu(note_emu) if note_emu else SH - Emu(800000),
             Emu(10000000), Emu(260000), [[(note, 12, CBLUE, False)]])

    # ---------- 3 学习目标 ----------
    def goals(self, items):
        s = self.add()
        self._header(s, "学习目标")
        cols = [(("知识", BLUE), items[0]), (("能力", ORANGE), items[1]),
                (("素质", CBLUE), items[2])]
        x = Emu(700000)
        for ((tag, _col), content) in cols:
            box(s, x, Emu(2200000), Emu(3300000), Emu(3400000),
                fill=RGBColor(0xF3, 0xF6, 0xFD))
            box(s, x, Emu(2200000), Emu(3300000), Emu(700000), fill=NAVY)
            text(s, x, Emu(2280000), Emu(3300000), Emu(600000),
                 [[(tag + "目标", 18, WHITE, True)]], PP_ALIGN.CENTER)
            text(s, x + Emu(180000), Emu(3100000), Emu(2940000), Emu(2300000),
                 [[(content, 13, NAVY, False)]], line_spacing=1.25)
            x += Emu(3550000)

    # ---------- 4 课前资讯回顾 ----------
    def review(self, question, check_item):
        s = self.add()
        self._header(s, "课前资讯回顾（六步之\u201c资讯·计划\u201d）")
        box(s, Emu(800000), Emu(2100000), Emu(10600000), Emu(1700000),
            fill=RGBColor(0xF3, 0xF6, 0xFD))
        box(s, Emu(800000), Emu(2100000), Emu(120000), Emu(1700000), fill=ORANGE)
        text(s, Emu(1100000), Emu(2280000), Emu(10000000), Emu(1400000),
             [[("预习快问：", 15, ORANGE, True)],
              [(question, 16, NAVY, False)]])
        text(s, Emu(800000), Emu(4100000), Emu(10600000), Emu(1500000),
             [[("课前检查：", 15, BLUE, True)],
              [(check_item, 14, GRAY, False)]])

    # ---------- 5 任务情境 ----------
    def scenario(self, desc, dept=None, illus=None):
        s = self.add()
        self._header(s, "任务情境（集团项目案例）")
        if dept:
            box(s, Emu(800000), Emu(1780000), Emu(10600000), Emu(700000),
                fill=NAVY)
            text(s, Emu(1000000), Emu(1880000), Emu(10200000), Emu(520000),
                 [[("◆ 重庆两江智造集团 · " + dept, 14, RGBColor(0xFF, 0xFF, 0xFF), True)]])
            if illus:
                box(s, Emu(800000), Emu(2580000), Emu(6300000), Emu(3600000),
                    fill=RGBColor(0xF7, 0xF9, 0xFD), line=BLUE)
                text(s, Emu(1000000), Emu(2740000), Emu(5900000), Emu(3300000),
                     [[(desc, 14.5, NAVY, False)]], line_spacing=1.3)
                try:
                    s.shapes.add_picture(
                        illus, Emu(7250000), Emu(2580000),
                        Emu(4150000), Emu(3600000))
                except Exception:
                    pass
            else:
                box(s, Emu(800000), Emu(2580000), Emu(10600000), Emu(3300000),
                    fill=RGBColor(0xF7, 0xF9, 0xFD), line=BLUE)
                text(s, Emu(1050000), Emu(2740000), Emu(10100000), Emu(3000000),
                     [[(desc, 15, NAVY, False)]], line_spacing=1.35)
        else:
            box(s, Emu(800000), Emu(2000000), Emu(10600000), Emu(3600000),
                fill=RGBColor(0xF7, 0xF9, 0xFD), line=BLUE)
            text(s, Emu(1050000), Emu(2200000), Emu(10100000), Emu(3300000),
                 [[(desc, 15, NAVY, False)]], line_spacing=1.35)

    # ---------- 6..n 内容页 ----------
    def content(self, idx, title, bullets):
        s = self.add()
        self._header(s, title, num=idx)
        y = Emu(2050000)
        for b in bullets:
            box(s, Emu(820000), y + Emu(110000), Emu(120000), Emu(120000),
                fill=ORANGE)
            text(s, Emu(1100000), y, Emu(10200000),
                 Emu(min(900000, 300000 + 240000 * (len(b) // 40 + 1))),
                 [[(b, 15, NAVY, False)]], line_spacing=1.25)
            y += Emu(min(1000000, 420000 + 260000 * (len(b) // 38 + 1)))

    # ---------- 决策矩阵 / 四步法 ----------
    def decision(self, title, header, rows, note="", y0_emu=None, colw=None, rowh=None):
        s = self.add()
        self._header(s, title)
        x0, y0 = Emu(800000), Emu(y0_emu if y0_emu is not None else 2000000)
        colw = colw or [Emu(3400000)] + [Emu(2400000)] * (len(header) - 1)
        rowh = rowh if rowh is not None else 620000
        # header row
        x = x0
        for htxt, w in zip(header, colw):
            box(s, x, y0, w, Emu(700000), fill=NAVY)
            text(s, x, y0 + Emu(120000), w, Emu(500000),
                 [[(htxt, 14, WHITE, True)]], PP_ALIGN.CENTER)
            x += w
        y = y0 + Emu(700000)
        for ri, row in enumerate(rows):
            x = x0
            for ci, val in enumerate(row):
                box(s, x, y, colw[ci], Emu(rowh),
                    fill=WHITE if ri % 2 else RGBColor(0xF3, 0xF6, 0xFD),
                    line=RGBColor(0xC9, 0xD4, 0xEE))
                text(s, x, y + Emu(max(100000, (rowh - 480000) // 2)), colw[ci], Emu(480000),
                     [[(str(val), 13, NAVY, ci == 0)]], PP_ALIGN.CENTER)
                x += colw[ci]
            y += Emu(rowh)
        if note:
            text(s, x0, y + Emu(120000), Emu(10600000), Emu(380000),
                 [[(note, 13, ORANGE, True)]])

    # ---------- 实施工单 ----------
    def task(self, wono, name, steps, std_pass, std_good):
        s = self.add()
        self._header(s, f"实施任务：工单{wono}《{name}》")
        box(s, Emu(800000), Emu(1950000), Emu(6300000), Emu(3900000),
            fill=RGBColor(0xF3, 0xF6, 0xFD))
        text(s, Emu(1000000), Emu(2100000), Emu(5900000), Emu(3600000),
             [[("实施步骤", 15, BLUE, True)]] +
             [[(f"{i+1}. {st}", 14, NAVY, False)] for i, st in enumerate(steps)],
             line_spacing=1.3)
        box(s, Emu(7400000), Emu(1950000), Emu(4000000), Emu(1850000),
            fill=RGBColor(0xE8, 0xF1, 0xE8), line=CBLUE)
        text(s, Emu(7550000), Emu(2080000), Emu(3750000), Emu(1650000),
             [[("及格标准（D）", 14, CBLUE, True)], [(std_pass, 13, NAVY, False)]])
        box(s, Emu(7400000), Emu(3950000), Emu(4000000), Emu(1850000),
            fill=RGBColor(0xFD, 0xF0, 0xE6), line=ORANGE)
        text(s, Emu(7550000), Emu(4080000), Emu(3750000), Emu(1650000),
             [[("优秀标准（B+）", 14, ORANGE, True)], [(std_good, 13, NAVY, False)]])

    # ---------- 三层检查 ----------
    def check(self):
        s = self.add()
        self._header(s, "过程检查（三层质检）")
        items = [("自检", "对照工单checklist逐项自查并截图留证"),
                 ("互检", "组间交换检查，互检表记录问题≥1条"),
                 ("师查", "教师抽查复核，重点盯D层完成度")]
        x = Emu(700000)
        for i, (t, d) in enumerate(items):
            box(s, x, Emu(2300000), Emu(3300000), Emu(2600000),
                fill=WHITE, line=BLUE)
            text(s, x, Emu(2450000), Emu(3300000), Emu(600000),
                 [[(f"{i+1}  {t}", 20, BLUE, True)]], PP_ALIGN.CENTER)
            text(s, x + Emu(200000), Emu(3200000), Emu(2900000), Emu(1500000),
                 [[(d, 13, GRAY, False)]], PP_ALIGN.CENTER, line_spacing=1.3)
            x += Emu(3550000)
        text(s, Emu(700000), Emu(5200000), Emu(10800000), Emu(500000),
             [[("复盘追问：如果当初选另一个方案，结果会怎样？（决策追责）",
                14, ORANGE, True)]], PP_ALIGN.CENTER)

    # ---------- 评价与课后 ----------
    def closing_tasks(self, homework, next_prev):
        s = self.add()
        self._header(s, "总结评价与课后任务")
        box(s, Emu(800000), Emu(1950000), Emu(10600000), Emu(1500000),
            fill=RGBColor(0xF3, 0xF6, 0xFD))
        text(s, Emu(1000000), Emu(2050000), Emu(10300000), Emu(1400000),
             [[("评价（10分钟）：", 15, BLUE, True)],
              [("按D/C/B/A量表当堂评分；工单24小时内反馈评语与改进建议。", 14, NAVY, False)]])
        y = Emu(3700000)
        for i, hw in enumerate(homework[:3]):
            text(s, Emu(850000), y, Emu(10600000), Emu(700000),
                 [[(f"{i+1}. ", 14, ORANGE, True), (hw, 14, NAVY, False)]],
                 line_spacing=1.2)
            y += Emu(720000)
        text(s, Emu(850000), y + Emu(100000), Emu(10600000), Emu(500000),
             [[("下次课预告：", 13, CBLUE, True), (next_prev, 13, GRAY, False)]])

    # ---------- 结束页 ----------
    def thanks(self, slogan=("以生为本", "做中学 · 学中做")):
        s = self.add()
        box(s, 0, 0, SW, SH, NAVY)
        box(s, 0, SH * 42 // 100, SW, Emu(50000), ORANGE)
        text(s, 0, SH * 24 // 100, SW, Emu(1500000),
             [[("THANKS", 54, WHITE, True)]], PP_ALIGN.CENTER)
        text(s, 0, SH * 50 // 100, SW, Emu(800000),
             [[("感谢聆听 · 动手实践出真知", 22, ORANGE, True)]], PP_ALIGN.CENTER)
        text(s, 0, SH * 66 // 100, SW, Emu(600000),
             [[(slogan[0] + "   |   " + slogan[1], 14,
                RGBColor(0xBB, 0xC6, 0xE6), False)]], PP_ALIGN.CENTER)

    # ---------- 内容页页眉 ----------
    def _header(self, s, title, num=None):
        box(s, 0, 0, SW, Emu(1150000), NAVY)
        box(s, 0, Emu(1150000), Emu(2600000), Emu(70000), ORANGE)
        if num:
            box(s, Emu(600000), Emu(230000), Emu(700000), Emu(700000), fill=ORANGE)
            text(s, Emu(600000), Emu(255000), Emu(700000), Emu(660000),
                 [[(str(num), 24, WHITE, True)]], PP_ALIGN.CENTER)
            text(s, Emu(1500000), Emu(300000), SW - Emu(2200000), Emu(700000),
                 [[(title, 24, WHITE, True)]])
        else:
            text(s, Emu(700000), Emu(300000), SW - Emu(1400000), Emu(700000),
                 [[(title, 24, WHITE, True)]])

    # ---------- v2: 知识详解页 ----------
    def knowledge(self, idx, points):
        s = self.add()
        self._header(s, "知识详解（讲）")
        y = Emu(1900000)
        for i, (t, d) in enumerate(points):
            box(s, Emu(760000), y, Emu(10900000), Emu(2100000),
                fill=RGBColor(0xF3, 0xF6, 0xFD))
            box(s, Emu(760000), y, Emu(130000), Emu(2100000), fill=BLUE)
            text(s, Emu(1050000), y + Emu(120000), Emu(2600000), Emu(600000),
                 [[(f"知识点{idx*2 + i + 1}  {t}", 17, NAVY, True)]])
            text(s, Emu(1050000), y + Emu(720000), Emu(10200000), Emu(1300000),
                 [[(d, 14, GRAY, False)]], line_spacing=1.3)
            y += Emu(2300000)

    # ---------- v2: 实施步骤与注意事项 ----------
    def steps_caution(self, steps, cautions):
        s = self.add()
        self._header(s, "实施步骤与注意事项（做）")
        box(s, Emu(700000), Emu(1850000), Emu(6300000), Emu(4100000),
            fill=WHITE, line=BLUE)
        text(s, Emu(900000), Emu(1980000), Emu(5900000), Emu(500000),
             [[("实施步骤", 17, BLUE, True)]])
        y = Emu(2600000)
        for i, st in enumerate(steps):
            box(s, Emu(920000), y + Emu(60000), Emu(480000), Emu(480000),
                fill=BLUE)
            text(s, Emu(920000), y + Emu(80000), Emu(480000), Emu(440000),
                 [[(str(i + 1), 15, WHITE, True)]], PP_ALIGN.CENTER)
            text(s, Emu(1520000), y, Emu(5300000), Emu(700000),
                 [[(st, 13.5, NAVY, False)]], line_spacing=1.15)
            y += Emu(min(760000, 330000 + 200000 * (len(st) // 30)))
        box(s, Emu(7350000), Emu(1850000), Emu(4300000), Emu(4100000),
            fill=RGBColor(0xFD, 0xF0, 0xE6), line=ORANGE)
        text(s, Emu(7550000), Emu(1980000), Emu(3950000), Emu(500000),
             [[("⚠ 注意事项", 17, ORANGE, True)]])
        text(s, Emu(7550000), Emu(2600000), Emu(3950000), Emu(3200000),
             [[(f"{i+1}. {c}", 13, NAVY, False)] for i, c in enumerate(cautions)],
             line_spacing=1.3, space_after=8)

    # ---------- v2: ABCD分层工单页 ----------
    def tiers(self, wono, sp, sg, extA):
        s = self.add()
        self._header(s, f"分层工单：工单{wono}（ABCD四级任务）")
        rows = [
            ("A 师傅", ORANGE, "进阶拓展层",
             extA + "；并担任组内示范讲解（以教代学），答辩加分"),
            ("B 组长", BLUE, "标准达成层",
             f"独立完成全部任务，达优秀标准：{sg}；组织组内互检与进度把控"),
            ("C 操作员", CBLUE, "基础完成层",
             f"在starter提示下完成核心任务，达及格标准：{sp}；如实记录过程数据"),
            ("D 学徒", GRAY, "跟做保底层",
             "跟做示例完成TODO-1最小任务，程序能运行出结果即达标；观摩A类示范并复述要点"),
        ]
        y = Emu(1800000)
        for tag, color, layer, desc in rows:
            box(s, Emu(700000), y, Emu(10900000), Emu(1050000), fill=WHITE,
                line=color)
            box(s, Emu(700000), y, Emu(1900000), Emu(1050000), fill=color)
            text(s, Emu(700000), y + Emu(130000), Emu(1900000), Emu(500000),
                 [[(tag, 17, WHITE, True)]], PP_ALIGN.CENTER)
            text(s, Emu(720000), y + Emu(640000), Emu(1900000), Emu(350000),
                 [[(layer, 11.5, WHITE, False)]], PP_ALIGN.CENTER)
            text(s, Emu(2800000), y + Emu(130000), Emu(8700000), Emu(850000),
                 [[(desc, 13, NAVY, False)]], line_spacing=1.2)
            y += Emu(1150000)
        text(s, Emu(700000), y + Emu(40000), Emu(10900000), Emu(400000),
             [[("分层规则：按学习意愿与能力定层，每4周按工单成绩动态升降（连续2次≥B升一层）", 12, GRAY, False)]])

    # ---------- v2: 课堂截图占位页 ----------
    def shot_placeholder(self):
        s = self.add()
        self._header(s, "课堂实况与运行截图（教师课后替换）")
        xs = [Emu(700000), Emu(5560000)]
        ys = [Emu(1900000), Emu(3700000)]
        for x in xs:
            for y in ys:
                box(s, x, y, Emu(4600000), Emu(1600000),
                    fill=RGBColor(0xF3, 0xF6, 0xFD), line=RGBColor(0xC9, 0xD4, 0xEE))
                text(s, x, y + Emu(620000), Emu(4600000), Emu(500000),
                     [[("［插入运行截图 / 实拍照片］", 14, GRAY, False)]],
                     PP_ALIGN.CENTER)

    def note(self, text):
        """演讲者备注：写入刚创建的那页幻灯片（Presenter View可见）"""
        idx = len(self.prs.slides._sldIdLst) - 1
        self.prs.slides[idx].notes_slide.notes_text_frame.text = text

    def save(self, path):
        self.prs.save(path)
