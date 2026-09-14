# -*- coding: utf-8 -*-
"""PPT v3页面扩展(继承Deck)：冲刺≈50页/课的深度讲义版页面"""
from ppt_gen import (BLUE, CBLUE, GRAY, NAVY, ORANGE, WHITE, Deck, box,
                     text)
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

LIGHT = WHITE.__class__(0xF3, 0xF6, 0xFD)
LINE = WHITE.__class__(0xC9, 0xD4, 0xEE)
CODE_BG = WHITE.__class__(0x0D, 0x11, 0x2A)
CODE_FG = WHITE.__class__(0xD6, 0xDE, 0xF6)
CODE_KEY = WHITE.__class__(0x7A, 0xB7, 0xFE)
ORANGE_L = WHITE.__class__(0xFD, 0xF0, 0xE6)
GREEN_L = WHITE.__class__(0xE8, 0xF1, 0xE8)
SW = Emu(12192000)


class DeckV3(Deck):

    # ---------- 章节过渡页 ----------
    # ---------- v9 本次课议程页（封面后，一眼看清六部分与分钟归属） ----------
    def agenda(self, items, tail=""):
        """items: [(编号, 名称, 分钟标注, 一句话)] ×6，2列×3行卡片"""
        s = self.add()
        self._header(s, "本次课议程（六部分 · 80分钟）")
        x0, y0 = Emu(760000), Emu(1850000)
        cw, ch, gx, gy = Emu(5150000), Emu(1400000), Emu(360000), Emu(220000)
        for i, (no, name, mins, note) in enumerate(items):
            x = x0 + (i % 2) * (cw + gx)
            y = y0 + (i // 2) * (ch + gy)
            box(s, x, y, cw, ch, LIGHT, line=BLUE)
            box(s, x, y, Emu(950000), ch, NAVY)
            text(s, x, y + Emu(380000), Emu(950000), Emu(600000),
                 [[(no, 22, WHITE, True)]], PP_ALIGN.CENTER)
            text(s, x + Emu(1120000), y + Emu(150000), Emu(3900000),
                 Emu(520000),
                 [[(name, 15, NAVY, True), ],
                  [(mins, 12, ORANGE, True)]])
            text(s, x + Emu(1120000), y + Emu(760000), Emu(3900000),
                 Emu(560000), [[(note, 11, GRAY, False)]], line_spacing=1.15)
        if tail:
            text(s, Emu(760000), Emu(6350000), Emu(10900000), Emu(400000),
                 [[(tail, 12, CBLUE, False)]])

    def section(self, no, title, sub=""):
        s = self.add()
        box(s, 0, 0, SW, Emu(6858000), NAVY)
        box(s, Emu(1000000), Emu(2650000), Emu(560000), Emu(1500000), ORANGE)
        text(s, Emu(1750000), Emu(2450000), Emu(9000000), Emu(1900000),
             [[(f"{no}", 66, ORANGE, True)],
              [(title, 32, WHITE, True)]])
        if sub:
            text(s, Emu(1750000), Emu(4600000), Emu(9000000), Emu(700000),
                 [[(sub, 16, WHITE.__class__(0xBB, 0xC6, 0xE6), False)]])
        # v5: 六部分进度点（当前章橙色高亮，其余暗点），全局导航
        try:
            cur = int(no)
        except (TypeError, ValueError):
            cur = 0
        dx = Emu(8100000)
        for i in range(1, 7):
            hot = (i == cur)
            box(s, dx, Emu(5900000), Emu(230000), Emu(230000),
                fill=ORANGE if hot else WHITE.__class__(0x2A, 0x35, 0x60),
                line=ORANGE if hot else None)
            text(s, dx - Emu(70000), Emu(6200000), Emu(440000), Emu(280000),
                 [[(f"0{i}", 9,
                    ORANGE if hot else WHITE.__class__(0x66, 0x74, 0xA8),
                    hot)]], PP_ALIGN.CENTER)
            dx += Emu(560000)

    # ---------- 课程定位页 ----------
    def positioning(self, cname, lines):
        """label左+value右同行式：5条×880000高不溢出（旧竖排第5条超出页面底部6858000）"""
        s = self.add()
        self._header(s, f"课程定位：{cname}（课程标准摘录）")
        y = Emu(1750000)
        for k, v in lines:
            box(s, Emu(760000), y, Emu(10900000), Emu(880000), LIGHT)
            box(s, Emu(760000), y, Emu(150000), Emu(880000), BLUE)
            text(s, Emu(1030000), y + Emu(240000), Emu(2300000), Emu(400000),
                 [[(k, 14, BLUE, True)]])
            text(s, Emu(3500000), y + Emu(90000), Emu(9100000), Emu(720000),
                 [[(v, 12, NAVY, False)]], line_spacing=1.15,
                 anchor=MSO_ANCHOR.MIDDLE)
            y += Emu(990000)

    # ---------- 知识图谱页 ----------
    def mindmap(self, center, nodes, title="本课知识图谱（先见森林，再见树木）"):
        """v2知识图谱：左侧主题枢纽 + 扇形连线 + 右侧五行分支卡（无重叠）"""
        s = self.add()
        self._header(s, title)
        # 左侧主题枢纽
        box(s, Emu(700000), Emu(2700000), Emu(1700000), Emu(1300000), ORANGE)
        text(s, Emu(760000), Emu(2900000), Emu(1580000), Emu(1100000),
             [[(center[:5], 17, WHITE, True)],
              [(center[5:10] if len(center) > 5 else "主题", 17, WHITE, True)]],
             PP_ALIGN.CENTER)
        text(s, Emu(700000), Emu(4300000), Emu(1700000), Emu(500000),
             [[(f"{len(nodes)}条分支导览", 11, GRAY, False)]], PP_ALIGN.CENTER)
        # 右侧分支卡（单列5行， chip+条目同行）
        from pptx.enum.shapes import MSO_CONNECTOR
        y = Emu(1650000)
        for lv1, subs in nodes:
            h = Emu(780000)
            cy = y + Emu(390000)
            # 枢纽→卡片 扇形连线
            ln = s.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, Emu(2400000), Emu(3350000),
                Emu(3450000), cy)
            ln.line.color.rgb = WHITE.__class__(0xB9, 0xC4, 0xDC)
            ln.line.width = Pt(1.6)
            box(s, Emu(3500000), y, Emu(8100000), h, LIGHT, line=BLUE)
            box(s, Emu(3500000), y, Emu(1500000), h, BLUE)
            text(s, Emu(3560000), y + Emu(220000), Emu(1380000), Emu(400000),
                 [[(lv1, 13, WHITE, True)]], PP_ALIGN.CENTER)
            text(s, Emu(5140000), y + Emu(90000), Emu(6340000), Emu(620000),
                 [[(" · ".join(subs), 11, NAVY, False)]],
                 line_spacing=1.15)
            y = y + Emu(920000)

    # ---------- 基础知识清单页(每页3条卡片) ----------
    def basics(self, idx, items):
        s = self.add()
        self._header(s, f"基础知识清单（{idx}）")
        y = Emu(1800000)
        for i, it in enumerate(items):
            n = (idx - 1) * 3 + i + 1
            box(s, Emu(760000), y, Emu(10900000), Emu(1350000), LIGHT)
            box(s, Emu(760000), y, Emu(1100000), Emu(1350000), NAVY)
            text(s, Emu(760000), y + Emu(330000), Emu(1100000), Emu(700000),
                 [[(f"B{n:02d}", 20, WHITE, True)]], PP_ALIGN.CENTER)
            text(s, Emu(2050000), y + Emu(380000), Emu(9400000), Emu(700000),
                 [[(it, 16, NAVY, False)]], line_spacing=1.2)
            y += Emu(1500000)

    # ---------- 技能点页 ----------
    def skills(self, items):
        s = self.add()
        self._header(s, "技能点：做完你应该会")
        y = Emu(1800000)
        for i, it in enumerate(items):
            box(s, Emu(760000), y, Emu(10900000), Emu(1000000), WHITE, line=ORANGE)
            box(s, Emu(820000), y + Emu(200000), Emu(600000), Emu(600000), ORANGE)
            text(s, Emu(820000), y + Emu(250000), Emu(600000), Emu(500000),
                 [[(f"S{i+1}", 15, WHITE, True)]], PP_ALIGN.CENTER)
            text(s, Emu(1650000), y + Emu(260000), Emu(9800000), Emu(600000),
                 [[(it, 15, NAVY, False)]])
            y += Emu(1130000)

    # ---------- 单层任务详情页 ----------
    def tier_detail(self, tier, goal, steps, std, extA=None):
        cfg = {"A": (ORANGE, "师傅 · 进阶拓展层"), "B": (BLUE, "组长 · 标准达成层"),
               "C": (CBLUE, "操作员 · 基础完成层"), "D": (GRAY, "学徒 · 跟做保底层")}
        color, label = cfg[tier]
        s = self.add()
        self._header(s, f"分层工单 · {tier} 层任务卡")
        box(s, Emu(760000), Emu(1750000), Emu(3200000), Emu(900000), color)
        text(s, Emu(760000), Emu(1900000), Emu(3200000), Emu(600000),
             [[(label, 17, WHITE, True)]], PP_ALIGN.CENTER)
        box(s, Emu(4300000), Emu(1750000), Emu(7300000), Emu(900000), LIGHT)
        text(s, Emu(4480000), Emu(1830000), Emu(7000000), Emu(760000),
             [[("目标: " + goal, 13.5, NAVY, True)]], line_spacing=1.2)
        if len(steps) >= 4:
            # v2: 4步及以上改两列，保证达标框不越页脚
            half = (len(steps) + 1) // 2
            for col, chunk in enumerate((steps[:half], steps[half:])):
                x = Emu(820000 + col * 5450000)
                yy = Emu(2850000)
                for j, st in enumerate(chunk):
                    no = col * half + j + 1
                    box(s, x, yy + Emu(40000), Emu(620000), Emu(620000), color)
                    text(s, x, yy + Emu(100000), Emu(620000), Emu(500000),
                         [[(str(no), 16, WHITE, True)]], PP_ALIGN.CENTER)
                    text(s, x + Emu(760000), yy + Emu(40000), Emu(4520000),
                         Emu(840000), [[(st, 12.5, NAVY, False)]],
                         line_spacing=1.15)
                    yy += Emu(900000)
            y = Emu(2850000 + 2 * 900000)
        else:
            y = Emu(2850000)
            for i, st in enumerate(steps):
                box(s, Emu(820000), y + Emu(40000), Emu(620000), Emu(620000), color)
                text(s, Emu(820000), y + Emu(100000), Emu(620000), Emu(500000),
                     [[(str(i + 1), 16, WHITE, True)]], PP_ALIGN.CENTER)
                text(s, Emu(1620000), y + Emu(70000), Emu(10000000), Emu(800000),
                     [[(st, 14, NAVY, False)]], line_spacing=1.2)
                y += Emu(890000)
        # 达标框（extA并入框内第二行，避免与页脚重叠/越界）
        box_h = Emu(980000) if extA else Emu(820000)
        box(s, Emu(760000), y + Emu(100000), Emu(10900000), box_h,
            ORANGE_L if tier in "AB" else GREEN_L, line=color)
        rows = [[("✔ 达标标准: " + std, 13.5, NAVY, True)]]
        if extA:
            rows.append([("(A层加码: " + extA + ")", 11.5, GRAY, False)])
        text(s, Emu(940000), y + Emu(200000), Emu(10500000),
             Emu(box_h - Emu(300000)), rows, line_spacing=1.15)

    # ---------- 代码页 ----------
    def code(self, title, code):
        s = self.add()
        self._header(s, f"关键代码：{title}")
        box(s, Emu(760000), Emu(1750000), Emu(10900000), Emu(4600000), CODE_BG)
        box(s, Emu(760000), Emu(1750000), Emu(10900000), Emu(120000),
            ORANGE)
        lines = code.splitlines()
        show = lines[:26]
        tb = text(s, Emu(1000000), Emu(2000000), Emu(10300000), Emu(4200000),
                  [[(ln if ln else " ", 12.5, CODE_FG, False)] for ln in show],
                  space_after=0, line_spacing=1.12)
        if len(lines) > 26:
            text(s, Emu(1000000), Emu(5900000), Emu(10300000), Emu(400000),
                 [[(f"... (共{len(lines)}行, 完整见starter/工单)", 11, GRAY, False)]])

    # ---------- 常见错误页 ----------
    def errors(self, rows):
        s = self.add()
        self._header(s, "常见错误速查（避坑清单）")
        y = Emu(1800000)
        for a, b, c in rows:
            box(s, Emu(760000), y, Emu(10900000), Emu(1180000), WHITE, line=LINE)
            box(s, Emu(760000), y, Emu(3000000), Emu(1180000), ORANGE_L)
            text(s, Emu(820000), y + Emu(100000), Emu(2800000), Emu(980000),
                 [[("⚠ " + a, 13, ORANGE, True)]], line_spacing=1.15)
            text(s, Emu(3950000), y + Emu(120000), Emu(2400000), Emu(940000),
                 [[("原因: " + b, 12, NAVY, False)]], line_spacing=1.15)
            text(s, Emu(6600000), y + Emu(120000), Emu(4900000), Emu(940000),
                 [[("✔ 解决: " + c, 12, NAVY, False)]], line_spacing=1.15)
            y += Emu(1300000)

    # ---------- 示意流程图页(程序绘制) ----------
    def fig_flow(self, title, steps, note=""):
        s = self.add()
        self._header(s, title)
        n = len(steps)
        w = Emu(min(3000000, 10800000 // max(n, 1) - 200000))
        gap = Emu(200000)
        total = w * n + gap * (n - 1)
        x = (SW - total) // 2
        y = Emu(2600000)
        for i, st in enumerate(steps):
            color = [BLUE, ORANGE, CBLUE, GRAY, BLUE, ORANGE][i % 6]
            box(s, x, y, w, Emu(1500000), color)
            text(s, x + Emu(60000), y + Emu(120000), w - Emu(120000),
                 Emu(1300000), [[(st, 13, WHITE, True)]],
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, line_spacing=1.15)
            if i < n - 1:
                text(s, x + w, y + Emu(520000), gap, Emu(500000),
                     [[("→", 18, ORANGE, True)]], PP_ALIGN.CENTER)
            x += w + gap
        if note:
            text(s, Emu(760000), Emu(4600000), Emu(10900000), Emu(900000),
                 [[(note, 13, GRAY, False)]], PP_ALIGN.CENTER)

    # ---------- 架构分层图 ----------
    def fig_arch(self, title, layers):
        """layers: [(层名, [组件...])] 自上而下分层（5层×960000不溢出）"""
        s = self.add()
        self._header(s, title)
        y = Emu(1800000)
        colors = [NAVY, BLUE, CBLUE, GRAY]
        for li, (lname, comps) in enumerate(layers):
            color = colors[li % 4]
            box(s, Emu(760000), y, Emu(2200000), Emu(880000), color)
            text(s, Emu(760000), y + Emu(240000), Emu(2200000), Emu(500000),
                 [[(lname, 14, WHITE, True)]], PP_ALIGN.CENTER)
            cw = Emu(10400000 // max(len(comps), 1) - 120000)
            for ci, c in enumerate(comps):
                cx = Emu(3200000) + ci * (cw + Emu(120000))
                box(s, cx, y + Emu(60000), cw, Emu(760000), LIGHT, line=color)
                text(s, cx + Emu(40000), y + Emu(120000), cw - Emu(80000),
                     Emu(680000), [[(c, 11.5, NAVY, False)]],
                     PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, line_spacing=1.1)
            y += Emu(960000)
        box(s, Emu(3200000), Emu(1800000), Emu(20000), y - Emu(1900000), LINE)

    # ---------- 对比柱图(手绘) ----------
    def fig_bars(self, title, series, ymax=None, note=""):
        """series: [(名称, 值, 颜色)]"""
        s = self.add()
        self._header(s, title)
        base_y = Emu(5400000)
        maxv = ymax or max(v for _, v, _ in series)
        n = len(series)
        bw = Emu(1400000)
        x0 = (SW - (bw + Emu(700000)) * n) // 2
        for i, (name, v, color) in enumerate(series):
            h = Emu(int(2800000 * v / maxv))
            x = x0 + i * (bw + Emu(700000))
            box(s, x, base_y - h, bw, h, color)
            text(s, x - Emu(200000), base_y - h - Emu(560000), bw + Emu(400000),
                 Emu(500000), [[(str(v), 15, NAVY, True)]], PP_ALIGN.CENTER)
            text(s, x - Emu(200000), base_y + Emu(80000), bw + Emu(400000),
                 Emu(500000), [[(name, 12, GRAY, False)]], PP_ALIGN.CENTER)
        box(s, Emu(1400000), base_y, Emu(9400000), Emu(30000), NAVY)
        if note:
            text(s, Emu(760000), Emu(6100000), Emu(10900000), Emu(500000),
                 [[(note, 12, CBLUE, False)]], PP_ALIGN.CENTER)

    # ---------- 课堂即时练习页 ----------
    def practice(self, no, question, hint):
        s = self.add()
        self._header(s, f"课堂即时练习 {no}")
        box(s, Emu(760000), Emu(2000000), Emu(10900000), Emu(1900000), LIGHT)
        box(s, Emu(760000), Emu(2000000), Emu(150000), Emu(1900000), ORANGE)
        text(s, Emu(1030000), Emu(2150000), Emu(10400000), Emu(1600000),
             [[("题目", 13, ORANGE, True)], [(question, 15, NAVY, False)]],
             line_spacing=1.3)
        text(s, Emu(760000), Emu(4200000), Emu(10900000), Emu(900000),
             [[("提示: " + hint, 12.5, GRAY, False)]])
        text(s, Emu(760000), Emu(5300000), Emu(10900000), Emu(500000),
             [[("⏱ 3分钟 · 同桌互查 · 举手验证", 12, CBLUE, False)]])

    # ---------- 提交物与评分细则页 ----------
    def submission(self, wono):
        s = self.add()
        self._header(s, f"提交物与评分（工单{wono}）")
        box(s, Emu(760000), Emu(1800000), Emu(5300000), Emu(3800000), LIGHT)
        text(s, Emu(960000), Emu(1950000), Emu(4900000), Emu(3500000),
             [[("提交清单", 15, BLUE, True)],
              [("1. 程序/notebook（含运行输出）", 13, NAVY, False)],
              [("2. 分层工单纸件（勾层+自评）", 13, NAVY, False)],
              [("3. 决策矩阵照片", 13, NAVY, False)],
              [("4. 过程记录/避坑笔记", 13, NAVY, False)],
              [("命名: 学号_姓名_" + wono, 12, GRAY, False)]], space_after=8)
        box(s, Emu(6350000), Emu(1800000), Emu(5300000), Emu(3800000), WHITE,
            line=ORANGE)
        text(s, Emu(6550000), Emu(1950000), Emu(4900000), Emu(3500000),
             [[("评分构成", 15, ORANGE, True)],
              [("功能完成 40%", 13, NAVY, False)],
              [("技术质量 25%（指标达标）", 13, NAVY, False)],
              [("代码规范 20%（PEP 8）", 13, NAVY, False)],
              [("工程素养 15%（留痕/协作）", 13, NAVY, False)],
              [("D≥60 / C≥70 / B≥80 / A≥90", 12.5, GRAY, True)]], space_after=8)

    # ---------- 80分钟个人节奏页 ----------
    def pace(self, tier="C"):
        s = self.add()
        self._header(s, "80分钟课堂节奏（学生自查）")
        rows = [("0-10", "资讯回顾+决策矩阵", "跟上讨论, 记下结论"),
                ("10-40", "第一轮实施", "完成本层核心任务主体"),
                ("40-45", "三层检查", "自检表勾完+互检1条"),
                ("45-70", "第二轮实施", "优化/冲标准/补漏"),
                ("70-80", "评价+课后布置", "自评打分, 记录课后任务")]
        y = Emu(1850000)
        for ri, (t, act, selfcheck) in enumerate(rows):
            box(s, Emu(760000), y, Emu(10900000), Emu(820000),
                WHITE if ri % 2 else LIGHT)
            box(s, Emu(760000), y, Emu(1800000), Emu(820000), NAVY)
            text(s, Emu(760000), y + Emu(180000), Emu(1800000), Emu(500000),
                 [[(t + " min", 14, WHITE, True)]], PP_ALIGN.CENTER)
            text(s, Emu(2800000), y + Emu(180000), Emu(3300000), Emu(500000),
                 [[(act, 13, NAVY, True)]])
            text(s, Emu(6300000), y + Emu(180000), Emu(5200000), Emu(500000),
                 [[("我要做到: " + selfcheck, 12, GRAY, False)]])
            y += Emu(920000)

    # ---------- 单条知识大卡 ----------
    def basic_card(self, no, text_):
        s = self.add()
        self._header(s, f"基础知识 B{no:02d}")
        box(s, Emu(1400000), Emu(2300000), Emu(9400000), Emu(2600000), LIGHT)
        box(s, Emu(1400000), Emu(2300000), Emu(200000), Emu(2600000), BLUE)
        text(s, Emu(700000), Emu(1250000), Emu(1800000), Emu(900000),
             [[(f"B{no:02d}", 30, ORANGE, True)]])
        text(s, Emu(1900000), Emu(2900000), Emu(8600000), Emu(1700000),
             [[(text_, 20, NAVY, True)]], line_spacing=1.3)
        text(s, Emu(1900000), Emu(5300000), Emu(8600000), Emu(600000),
             [[("——用你自己的话, 向同桌复述一遍", 13, GRAY, False)]])

    # ---------- 知识快闪页 ----------
    def quick_check(self, items):
        s = self.add()
        self._header(s, "知识快闪：60秒复述以下要点")
        y = Emu(1800000)
        for i, it in enumerate(items, 1):
            text(s, Emu(800000), y, Emu(10800000), Emu(700000),
                 [[(f"{i}. ", 15, ORANGE, True), (it, 15, NAVY, False)]])
            y += Emu(760000)
        text(s, Emu(800000), y + Emu(120000), Emu(10800000), Emu(500000),
             [[("同桌互述 → 抽2人全班复述（费曼：能讲清才算懂）", 13, CBLUE, True)]])

    # ---------- 素养目标页 ----------
    def literacy(self, lines):
        s = self.add()
        self._header(s, "素养目标（课程思政融入点）")
        y = Emu(1900000)
        icons = ["责任感", "工匠精神", "数据伦理"]
        for i, ln in enumerate(lines):
            box(s, Emu(760000), y, Emu(10900000), Emu(1150000), LIGHT)
            box(s, Emu(760000), y, Emu(1900000), Emu(1150000), NAVY)
            text(s, Emu(760000), y + Emu(300000), Emu(1900000), Emu(600000),
                 [[(icons[i % 3], 15, WHITE, True)]], PP_ALIGN.CENTER)
            text(s, Emu(2900000), y + Emu(160000), Emu(8600000), Emu(900000),
                 [[(ln, 13.5, NAVY, False)]], line_spacing=1.25)
            y += Emu(1280000)

    # ---------- 下次课预告页 ----------
    def preview(self, title, quiz, week_no):
        s = self.add()
        self._header(s, f"下次课预告（第{week_no}周）")
        box(s, Emu(760000), Emu(1900000), Emu(10900000), Emu(1600000), LIGHT)
        box(s, Emu(760000), Emu(1900000), Emu(150000), Emu(1600000), ORANGE)
        text(s, Emu(1020000), Emu(2050000), Emu(10500000), Emu(1400000),
             [[("主题: " + title, 17, NAVY, True)],
              [("预习快问: " + quiz[:60], 13, GRAY, False)]], space_after=6)
        text(s, Emu(760000), Emu(3900000), Emu(10900000), Emu(900000),
             [[("课前完成：导学单 + 预习自测 + starter环境自检", 14, BLUE, True)],
              [("带着问题来上课——Predict先于Run", 12.5, GRAY, False)]],
             space_after=6)

    # ---------- 评估指标公式卡(AI通用) ----------
    def metrics_card(self):
        s = self.add()
        self._header(s, "评估指标公式卡（贴在显示器上）")
        cards = [("准确率 Accuracy", "对的 / 全部", "总体把握, 不均衡时会骗人"),
                 ("精确率 Precision", "报的真 / 全部报的", "误报成本高时看它"),
                 ("召回率 Recall", "抓到的真 / 全部真的", "漏报(漏检)代价高时看它"),
                 ("F1", "2PR / (P+R)", "精确与召回的调和平均")]
        pos = [(760000, 1850000), (5760000, 1850000),
               (760000, 3700000), (5760000, 3700000)]
        for (x, y), (name, formula, note) in zip(pos, cards):
            box(s, Emu(x), Emu(y), Emu(4900000), Emu(1650000), WHITE, line=BLUE)
            box(s, Emu(x), Emu(y), Emu(4900000), Emu(540000), BLUE)
            text(s, Emu(x + 150000), y + Emu(60000), Emu(4600000), Emu(450000),
                 [[(name, 14, WHITE, True)]])
            text(s, Emu(x + 150000), y + Emu(620000), Emu(4600000), Emu(500000),
                 [[(formula, 16, ORANGE, True)]])
            text(s, Emu(x + 150000), y + Emu(1150000), Emu(4600000), Emu(450000),
                 [[(note, 11.5, GRAY, False)]])


# ==== v4 全局装饰: 页脚+页码+集团标识 ====
_FOOT_CACHE = {"n": 0}


def _decorate(deck, label=""):
    """给全部页面加页脚(细线+课程名+页码)与集团小标识; 封面/谢谢页跳过"""
    from pptx.util import Emu as _E
    total = len(deck.prs.slides._sldIdLst)
    for idx, slide in enumerate(deck.prs.slides, 1):
        if idx == 1 or idx == total:
            continue
        # 底部细线
        box(slide, _E(500000), _E(6500000), _E(11100000), _E(20000),
            WHITE.__class__(0xC9, 0xD4, 0xEE))
        # 左下: 集团+课程
        text(slide, _E(500000), _E(6540000), _E(7000000), _E(300000),
             [[("两江智造集团 · " + label, 10,
                WHITE.__class__(0x8A, 0x97, 0xB5), False)]])
        # 右下: 页码
        text(slide, _E(10600000), _E(6540000), _E(1000000), _E(300000),
             [[(f"{idx:02d} / {total:02d}", 10,
                WHITE.__class__(0x8A, 0x97, 0xB5), False)]],
             __import__("pptx.enum.text", fromlist=["PP_ALIGN"]).PP_ALIGN.RIGHT)
